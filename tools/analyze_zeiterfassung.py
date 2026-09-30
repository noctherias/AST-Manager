"""Create a reproducible inventory of the Zeiterfassung macro workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


def value(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def workbook_inventory(source: Path) -> dict:
    raw = source.read_bytes()
    result: dict = {
        "file": {
            "name": source.name,
            "size": len(raw),
            "sha256": hashlib.sha256(raw).hexdigest(),
        }
    }

    with zipfile.ZipFile(source) as archive:
        members = archive.namelist()
        result["package"] = {
            "member_count": len(members),
            "members": members,
            "has_vba": "xl/vbaProject.bin" in members,
            "vba_sha256": hashlib.sha256(archive.read("xl/vbaProject.bin")).hexdigest()
            if "xl/vbaProject.bin" in members else None,
            "custom_ui": [name for name in members if name.lower().startswith("customui/")],
            "controls": [name for name in members if "control" in name.lower() or "active" in name.lower()],
            "drawings": [name for name in members if name.startswith("xl/drawings/")],
            "external_links": [name for name in members if name.startswith("xl/externalLinks/")],
        }
        macro_strings = []
        if "xl/vbaProject.bin" in members:
            blob = archive.read("xl/vbaProject.bin")
            for match in re.findall(rb"[ -~]{8,}", blob):
                text = match.decode("cp1252", errors="replace")
                if any(token in text.lower() for token in (
                    "sub ", "function ", "worksheet", "workbook", "button", "module", "export", "print"
                )):
                    macro_strings.append(text[:300])
        result["package"]["macro_strings"] = macro_strings[:200]

    wb = load_workbook(source, keep_vba=True, data_only=False, keep_links=True)
    cached = load_workbook(source, keep_vba=True, data_only=True, keep_links=True)
    defined_names = []
    for name, entry in wb.defined_names.items():
        try:
            destinations = list(entry.destinations)
        except Exception:
            destinations = []
        defined_names.append({
            "name": name,
            "type": entry.type,
            "attr_text": entry.attr_text,
            "hidden": entry.hidden,
            "local_sheet_id": entry.localSheetId,
            "destinations": destinations,
        })
    result["workbook"] = {
        "sheet_order": wb.sheetnames,
        "active_sheet": wb.active.title,
        "defined_names": defined_names,
        "calculation": {
            "mode": value(wb.calculation.calcMode),
            "full_calc_on_load": value(wb.calculation.fullCalcOnLoad),
            "force_full_calc": value(wb.calculation.forceFullCalc),
        },
        "properties": {
            "title": wb.properties.title,
            "subject": wb.properties.subject,
            "creator": wb.properties.creator,
            "description": wb.properties.description,
        },
    }

    sheets = []
    for ws in wb.worksheets:
        cached_ws = cached[ws.title]
        cells = []
        formulas = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None and not cell.has_style:
                    continue
                item = {
                    "coordinate": cell.coordinate,
                    "value": value(cell.value),
                    "cached": value(cached_ws[cell.coordinate].value),
                    "data_type": cell.data_type,
                    "number_format": cell.number_format,
                    "style_id": cell.style_id,
                    "locked": cell.protection.locked,
                    "hidden": cell.protection.hidden,
                }
                if cell.value is not None:
                    cells.append(item)
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    formulas.append(item)
        validations = []
        for validation in ws.data_validations.dataValidation:
            validations.append({
                "type": validation.type,
                "ranges": str(validation.sqref),
                "formula1": value(validation.formula1),
                "formula2": value(validation.formula2),
                "allow_blank": validation.allowBlank,
                "error": validation.error,
                "prompt": validation.prompt,
            })
        conditional = []
        for conditional_range, rules in ws.conditional_formatting._cf_rules.items():
            conditional.append({
                "range": str(conditional_range.sqref),
                "rules": [{
                    "type": rule.type,
                    "operator": value(rule.operator),
                    "formula": [value(v) for v in rule.formula],
                    "priority": rule.priority,
                } for rule in rules],
            })
        tables = []
        for table in ws.tables.values():
            tables.append({"name": table.name, "display_name": table.displayName, "range": table.ref})
        sheets.append({
            "title": ws.title,
            "state": ws.sheet_state,
            "dimensions": ws.calculate_dimension(),
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "freeze_panes": value(ws.freeze_panes),
            "auto_filter": ws.auto_filter.ref,
            "merged_cells": [str(rng) for rng in ws.merged_cells.ranges],
            "print_area": value(ws.print_area),
            "print_titles": value(ws.print_titles),
            "page_setup": {
                "orientation": ws.page_setup.orientation,
                "paper_size": ws.page_setup.paperSize,
                "fit_to_width": ws.page_setup.fitToWidth,
                "fit_to_height": ws.page_setup.fitToHeight,
                "scale": ws.page_setup.scale,
            },
            "margins": {
                "left": ws.page_margins.left,
                "right": ws.page_margins.right,
                "top": ws.page_margins.top,
                "bottom": ws.page_margins.bottom,
                "header": ws.page_margins.header,
                "footer": ws.page_margins.footer,
            },
            "protection": {
                "enabled": bool(ws.protection.sheet),
                "password": ws.protection.password,
            },
            "hidden_rows": [index for index, dim in ws.row_dimensions.items() if dim.hidden],
            "hidden_columns": [index for index, dim in ws.column_dimensions.items() if dim.hidden],
            "column_widths": {index: dim.width for index, dim in ws.column_dimensions.items() if dim.width is not None},
            "row_heights": {str(index): dim.height for index, dim in ws.row_dimensions.items() if dim.height is not None},
            "data_validations": validations,
            "conditional_formatting": conditional,
            "tables": tables,
            "charts": len(ws._charts),
            "images": len(ws._images),
            "formula_count": len(formulas),
            "formulas": formulas,
            "populated_cells": cells,
        })
    result["sheets"] = sheets
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    inventory = workbook_inventory(args.source)
    text = json.dumps(inventory, indent=2, ensure_ascii=False)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text + "\n", encoding="utf-8")
    print(text)


if __name__ == "__main__":
    main()
