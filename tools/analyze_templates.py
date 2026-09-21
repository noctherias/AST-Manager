"""Read-only, reproducible inventory of every supplied workbook and PDF form."""
from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]


def xml(value):
    try:
        return ET.tostring(value.to_tree(), encoding="unicode")
    except Exception:
        return str(value)


def analyze():
    result = {"workbooks": [], "pdf": {}}
    for path in sorted((ROOT / "templates").glob("*.xlsx")):
        wb = openpyxl.load_workbook(path, data_only=False, keep_links=True)
        cache = openpyxl.load_workbook(path, data_only=True)
        item = {
            "file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "defined_names": {n: xml(d) for n, d in wb.defined_names.items()},
            "calculation": xml(wb.calculation), "external_links": [xml(l) for l in wb._external_links],
            "sheets": [],
        }
        with ZipFile(path) as archive:
            item["package_parts"] = archive.namelist()
            item["workbook_xml"] = archive.read("xl/workbook.xml").decode()
            item["styles_xml"] = archive.read("xl/styles.xml").decode()
            item["embedded_logic_xml"] = {
                n: archive.read(n).decode() for n in archive.namelist()
                if ("externalLink" in n or "connections" in n or "vbaProject" in n)
                and n.endswith(".xml")
            }
        for ws in wb:
            sheet = {
                "name": ws.title, "state": ws.sheet_state, "dimensions": ws.calculate_dimension(),
                "print_area": str(ws.print_area), "print_title_rows": ws.print_title_rows,
                "print_title_cols": ws.print_title_cols, "page_setup": xml(ws.page_setup),
                "page_margins": xml(ws.page_margins), "print_options": xml(ws.print_options),
                "header_footer": xml(ws.HeaderFooter), "sheet_properties": xml(ws.sheet_properties),
                "views": xml(ws.views), "protection": xml(ws.protection),
                "freeze_panes": ws.freeze_panes, "auto_filter": xml(ws.auto_filter),
                "defined_names": {n: xml(d) for n, d in ws.defined_names.items()},
                "merged_cells": [str(r) for r in ws.merged_cells.ranges],
                "validations": [xml(d) for d in ws.data_validations.dataValidation],
                "conditional_formats": [{"range": str(k.sqref), "rules": [xml(r) for r in v]}
                                        for k, v in ws.conditional_formatting._cf_rules.items()],
                "tables": [xml(t) for t in ws.tables.values()],
                "row_dimensions": {str(k): xml(v) for k, v in ws.row_dimensions.items()},
                "column_dimensions": {str(k): xml(v) for k, v in ws.column_dimensions.items()},
                "cells": {}, "formulas": {}, "dependencies": {},
                "images": len(ws._images), "charts": len(ws._charts),
            }
            for row in ws:
                for c in row:
                    if c.value is None:
                        continue
                    cell = {"value": c.value, "type": c.data_type, "number_format": c.number_format,
                            "locked": c.protection.locked,
                            "comment": c.comment.text if c.comment else None,
                            "hyperlink": c.hyperlink.target if c.hyperlink else None}
                    if c.data_type == "f":
                        sheet["formulas"][c.coordinate] = c.value
                        cell["cached_value"] = cache[ws.title][c.coordinate].value
                        sheet["dependencies"][c.coordinate] = [
                            t.value for t in Tokenizer(c.value).items
                            if t.type == "OPERAND" and t.subtype == "RANGE"]
                    sheet["cells"][c.coordinate] = cell
            item["sheets"].append(sheet)
        result["workbooks"].append(item)
    path = ROOT / "templates" / "Vorlage_Lohnausweis.pdf"
    reader = PdfReader(path)
    result["pdf"] = {
        "file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "pages": len(reader.pages), "metadata": dict(reader.metadata or {}),
        "fields": {n: dict(v) for n, v in (reader.get_fields() or {}).items()},
        "widgets": [], "text": [p.extract_text() for p in reader.pages],
    }
    result["pdf"]["field_actions"] = {}
    for name, field in (reader.get_fields() or {}).items():
        for trigger, action in field.get("/AA", {}).items():
            script = action.get_object().get("/JS")
            if script is not None:
                script = script.get_object()
                result["pdf"]["field_actions"][name + ":" + trigger] = (
                    script.get_data().decode("latin1") if hasattr(script, "get_data") else str(script))
    for page_no, page in enumerate(reader.pages):
        for ref in page.get("/Annots", []):
            obj = ref.get_object()
            if obj.get("/Subtype") == "/Widget":
                parent = obj.get("/Parent")
                result["pdf"]["widgets"].append({
                    "page": page_no, "object_id": ref.idnum,
                    "name": obj.get("/T"), "type": obj.get("/FT"),
                    "parent": str(parent), "value": obj.get("/V"),
                    "rect": list(obj.get("/Rect", [])), "tooltip": obj.get("/TU"),
                    "appearance_states": list(obj.get("/AP", {}).get("/N", {}).keys())
                        if obj.get("/FT") == "/Btn" else [],
                    "actions": str(obj.get("/AA", "")),
                    "max_length": obj.get("/MaxLen"),
                })
    (ROOT / "docs" / "template_inventory.json").write_text(
        json.dumps(result, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    for wb in result["workbooks"]:
        print("\nFILE", wb["file"], "NAMES", wb["defined_names"], "LINKS", wb["external_links"])
        for sheet in wb["sheets"]:
            print("SHEET", sheet["name"], sheet["dimensions"], "PRINT", sheet["print_area"],
                  "NAMES", sheet["defined_names"], "DV", sheet["validations"])
            for cell, value in sheet["cells"].items():
                print(cell, repr(value["value"]), "CACHE", value.get("cached_value", ""))
    print("\nPDF FIELDS")
    for name, field in result["pdf"]["fields"].items():
        print(name, repr(field))
    print("PDF TEXT", result["pdf"]["text"])


if __name__ == "__main__":
    analyze()
