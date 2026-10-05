"""Read existing AST debtor and time-tracking workbooks into SQLite."""
from __future__ import annotations

from calendar import monthrange
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


def _date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"Ungültiges Datum: {text or 'leer'}")


def _cents(value):
    if value in (None, ""):
        return 0
    text = str(value).strip().replace("CHF", "").replace("'", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")
    try:
        return int((Decimal(text) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except InvalidOperation as exc:
        raise ValueError(f"Ungültiger Betrag: {value}") from exc


def import_debtors(db, path):
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    sheet = workbook.active
    header = None
    for row in range(1, min(sheet.max_row, 100) + 1):
        values = [str(sheet.cell(row, column).value or "").strip().casefold() for column in range(1, sheet.max_column + 1)]
        if "datum rechnung" in values and "nummer rechnung" in values:
            header = row
            columns = {value: index + 1 for index, value in enumerate(values)}
            break
    if not header:
        raise ValueError("Die Debitoren-Spalten wurden in dieser Excel-Datei nicht gefunden.")
    imported = skipped = payments = 0
    errors = []
    existing_numbers = {row["number"].casefold() for row in db.invoices()}
    customers = {row["name"].casefold(): row["id"] for row in db.customers()}
    for row in range(header + 1, sheet.max_row + 1):
        number = str(sheet.cell(row, columns["nummer rechnung"]).value or "").strip()
        name = str(sheet.cell(row, columns.get("name", 0)).value or "").strip() if columns.get("name") else ""
        if not number and not name:
            continue
        try:
            if not number or not name:
                raise ValueError("Rechnungsnummer oder Kundenname fehlt")
            if number.casefold() in existing_numbers:
                skipped += 1
                continue
            customer = customers.get(name.casefold())
            if not customer:
                customer = db.save_customer({"name": name})
                customers[name.casefold()] = customer
            issued = _date(sheet.cell(row, columns["datum rechnung"]).value)
            amount = _cents(sheet.cell(row, columns["betrag rechnung"]).value)
            paid = _cents(sheet.cell(row, columns["betrag bezahlt"]).value)
            if paid > amount:
                raise ValueError("Der bezahlte Betrag ist grösser als der Rechnungsbetrag")
            invoice = db.save_invoice({"number": number, "customer_id": customer, "issued": issued,
                                       "due": issued, "amount": amount, "note": "Aus Excel importiert"})
            if paid:
                valuta = _date(sheet.cell(row, columns["valuta datum"]).value)
                db.add_payment(invoice, valuta, paid, "Aus Excel importiert")
                payments += 1
            existing_numbers.add(number.casefold())
            imported += 1
        except Exception as exc:
            errors.append(f"Zeile {row}: {exc}")
    workbook.close()
    if errors and not imported:
        raise ValueError("Keine Rechnung importiert. " + errors[0])
    return {"imported": imported, "skipped": skipped, "payments": payments, "errors": errors}


def _minutes(value):
    if value in (None, "") or isinstance(value, str) and value.startswith("="):
        return None
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    if isinstance(value, timedelta):
        return round(value.total_seconds() / 60)
    if isinstance(value, (int, float, Decimal)):
        return round(float(value) * 1440)
    text = str(value).strip()
    try:
        hours, minutes = text.split(":", 1)
        return int(hours) * 60 + int(minutes[:2])
    except Exception as exc:
        raise ValueError(f"Ungültige Uhrzeit: {value}") from exc


def _decimal_work_minutes(value):
    if value in (None, ""):
        return 0
    if isinstance(value, (int, float, Decimal)):
        return round(float(value) * 60)
    try:
        return round(float(str(value).strip().replace(",", ".")) * 60)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Ungültige Arbeitszeit: {value}") from exc


def import_timesheet(db, path, employee_id, overwrite=False):
    workbook = load_workbook(Path(path), read_only=True, data_only=False, keep_vba=True)
    if "Voreinstellungen" not in workbook.sheetnames:
        raise ValueError("Diese Datei ist keine unterstützte AST-Zeiterfassung.")
    year = int(workbook["Voreinstellungen"]["C2"].value)
    month_sheets = workbook.worksheets[5:17]
    if len(month_sheets) != 12:
        raise ValueError("Die zwölf Monatsblätter wurden nicht gefunden.")
    records = []
    for month, sheet in enumerate(month_sheets, 1):
        hours_header = str(sheet["D3"].value or "").replace("\n", " ").strip().lower()
        direct_hours = hours_header in ("arbeitszeit", "arbeitszeit (h)")
        decimal_hours = hours_header == "arbeitszeit (h)"
        for day_number in range(1, monthrange(year, month)[1] + 1):
            row = day_number + 3
            values = [sheet.cell(row, column).value for column in (4, 5, 6, 7, 8, 10, 15)]
            if not any(value not in (None, "") and not (isinstance(value, str) and value.startswith("=")) for value in values):
                continue
            if direct_hours:
                start_1 = end_1 = start_2 = end_2 = None
                pause = 0
                effective_minutes = (_decimal_work_minutes(values[0]) if decimal_hours
                                     else (_minutes(values[0]) or 0))
            else:
                start_1, end_1, start_2, end_2 = (_minutes(value) for value in values[:4])
                pause = _minutes(values[4]) or 0
                effective_minutes = 0
                for start_value, end_value in ((start_1, end_1), (start_2, end_2)):
                    if start_value is not None and end_value is not None:
                        effective_minutes += (end_value - start_value) % 1440
                effective_minutes = max(0, effective_minutes - pause)
            code = "" if values[5] in (None, "") else str(values[5]).strip().upper()
            note = "" if values[6] in (None, "") else str(values[6]).strip()
            records.append({"employee_id": employee_id, "day": date(year, month, day_number).isoformat(),
                            "start_1": start_1, "end_1": end_1, "start_2": start_2, "end_2": end_2,
                            "break_minutes": pause, "worked_minutes": effective_minutes,
                            "code": code, "note": note})
    workbook.close()
    if not records:
        raise ValueError("In der Excel-Datei wurden keine erfassten Arbeitstage gefunden.")
    result = db.save_time_records(records, overwrite)
    result["year"] = year
    return result
