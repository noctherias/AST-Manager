"""Independent evaluation of every original Excel formula against native rules.

Only tests/tooling read xlsx. Production does not import openpyxl.
"""
import ast
from decimal import Decimal
from pathlib import Path
import random
import unittest

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries

from ast_app.domain import period_balance

ROOT = Path(__file__).resolve().parents[1]


class ReferenceEvaluator:
    def __init__(self, workbook, inputs):
        self.wb, self.inputs = workbook, inputs
        self.cache = {}

    def reference(self, token, sheet):
        if "!" in token:
            sheet, token = token.rsplit("!", 1)
            sheet = sheet.strip("'").replace("''", "'")
        token = token.replace("$", "")
        if ":" in token:
            c1, r1, c2, r2 = range_boundaries(token)
            ws = self.wb[sheet]
            return sum((self.cell(sheet, ws.cell(r, c).coordinate) for r in range(r1 or 1, (r2 or ws.max_row) + 1)
                        for c in range(c1 or 1, (c2 or ws.max_column) + 1)), Decimal(0))
        return self.cell(sheet, token)

    def cell(self, sheet, cell):
        key = (sheet, cell)
        if key in self.inputs:
            return Decimal(str(self.inputs[key]))
        if key in self.cache:
            return self.cache[key]
        v = self.wb[sheet][cell].value
        if isinstance(v, (int, float)):
            return Decimal(str(v))
        if not isinstance(v, str) or not v.startswith("="):
            return Decimal(0)
        expression = ""
        for tok in Tokenizer(v).items:
            if tok.type == "OPERAND" and tok.subtype == "RANGE":
                expression += str(self.reference(tok.value, sheet))
            elif tok.type == "FUNC" and tok.value == "SUM(":
                expression += "("
            elif tok.type in ("OPERATOR-INFIX", "OPERATOR-PREFIX", "OPERATOR-POSTFIX", "PAREN", "WHITE-SPACE") or tok.value == ")":
                expression += tok.value
            elif tok.type == "OPERAND" and tok.subtype == "NUMBER":
                expression += tok.value
            else:
                raise AssertionError(f"Uncovered formula token {tok.type}: {tok.value}")
        tree = ast.parse(expression, mode="eval")

        def evaluate(node):
            if isinstance(node, ast.Expression): return evaluate(node.body)
            if isinstance(node, ast.Constant): return Decimal(str(node.value))
            if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.UAdd): return evaluate(node.operand)
            if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.USub): return -evaluate(node.operand)
            if isinstance(node, ast.BinOp) and isinstance(node.op, ast.Add): return evaluate(node.left) + evaluate(node.right)
            if isinstance(node, ast.BinOp) and isinstance(node.op, ast.Sub): return evaluate(node.left) - evaluate(node.right)
            raise AssertionError("Unsupported reference expression")

        self.cache[key] = evaluate(tree)
        return self.cache[key]


class TemplateParityTests(unittest.TestCase):
    def test_every_hours_formula_with_inputs_and_cached_values(self):
        rng = random.Random(90120)
        formula_count = 0
        for filename, apprentice in [("Stundennachweis_Vorlage.xlsx", False), ("Stundennachweis_Lehrling_Vorlage.xlsx", True)]:
            wb = openpyxl.load_workbook(ROOT / "templates" / filename, data_only=False)
            cached = openpyxl.load_workbook(ROOT / "templates" / filename, data_only=True)
            source_formulas = {(s.title, c.coordinate) for s in wb for row in s for c in row if c.data_type == "f"}
            formula_count += len(source_formulas)
            for scenario in range(16):
                inputs, expected = {}, {}
                previous = None
                for idx, ws in enumerate(wb):
                    vacation_end = 21 if apprentice else 78 if idx == 2 else 65
                    total_row = 22 if apprentice else vacation_end + 1
                    sick_col, accident_col = ("L", "O") if idx == 0 else ("N", "Q")
                    values = {}
                    for kind, col, end in [("vacation", "B", vacation_end), ("overtime", "E", 21), ("sick", sick_col, 21), ("accident", accident_col, 21)]:
                        values[kind] = 0
                        for row in range(4, end + 1):
                            amount = rng.randint(-50 if kind == "overtime" else 0, 100) * 25 if scenario else 0
                            values[kind] += amount
                            inputs[(ws.title, f"{col}{row}")] = Decimal(amount) / 100
                    b = period_balance(21625, values["vacation"], values["overtime"], previous["balance"] if previous else 0,
                                       values["sick"], values["accident"], previous["sick_total"] if previous else 0,
                                       previous["accident_total"] if previous else 0)
                    mapping = {f"B{total_row}": b["remaining"], "E22": b["overtime"],
                               f"{sick_col}22": b["sick_total"], f"{accident_col}22": b["accident_total"]}
                    if idx == 0: mapping["G22"] = b["balance"]
                    else: mapping.update({"G22": b["carry"], "I22": b["balance"]})
                    if not apprentice:
                        mapping.update({"G45": b["remaining"], "G46": b["overtime"], "G47": b["balance"]})
                        if idx: mapping["G44"] = b["carry"]
                    expected.update({(ws.title, cell): Decimal(value)/100 for cell, value in mapping.items()})
                    previous = b
                self.assertEqual(set(expected), source_formulas)
                evaluator = ReferenceEvaluator(wb, inputs)
                for (sheet, cell), value in expected.items():
                    with self.subTest(file=filename, scenario=scenario, sheet=sheet, cell=cell):
                        self.assertEqual(evaluator.cell(sheet, cell), value)
                        if scenario == 0:
                            self.assertEqual(Decimal(str(cached[sheet][cell].value)), value)
        self.assertEqual(formula_count, 39)

    def test_debtor_sum_entire_column(self):
        wb = openpyxl.load_workbook(ROOT / "templates" / "Debitoren_Vorlage.xlsx")
        sheet = wb.active.title
        cells = {(sheet, "G14"): "100.25", (sheet, "G328"): "500", (sheet, "G590"): "80.30"}
        self.assertEqual(ReferenceEvaluator(wb, cells).cell(sheet, "K13"), Decimal("680.55"))


if __name__ == "__main__": unittest.main()
